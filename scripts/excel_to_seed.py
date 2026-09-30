#!/usr/bin/env python3
"""Régénère supabase/seed.sql + seed_data.json depuis VISION PATRIMOINE 030425.xlsx

Modèle Excel (onglets Amort Beneduc / Troika) :
  CRD(mois) = max(0, CRD_réf − mensualité × mois_depuis_début)
où CRD_réf = colonne E à la 1ʳᵉ échéance du tableau d’amort (nov. 2022 pour Beneduc/Troika).

Les prêts à taux (ex. Bd Libération 3,3 %, RP ~3,76 %) pourront être saisis en RATE_BASED
plus tard ; l’app génère alors le tableau via loanCalculator.
"""
import json, re, uuid
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    raise SystemExit("pip install openpyxl")

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "VISION PATRIMOINE 030425.xlsx"
OUT_JSON = ROOT / "supabase" / "seed_data.json"
OUT_SQL = ROOT / "supabase" / "seed.sql"

# 1ʳᵉ ligne des tableaux d’amort Excel (= date où la colonne E est vraie)
AMORT_START = {
    "beneduc": "2022-11-01",
    "troika": "2022-11-01",
    "lavista": "2023-05-01",
    "rp": "2024-11-01",
}

def uid(s: str) -> str:
    return str(uuid.uuid5(uuid.NAMESPACE_URL, f"vision:{s}"))

def parse_fin(v):
    if isinstance(v, datetime):
        y, m = v.year, v.month
        if y < 2000:
            y += 100
        return f"{y:04d}-{m:02d}-01"
    if isinstance(v, str):
        s = v.strip()
        ml = s.lower().replace("é", "e").replace("û", "u")
        month_map = [("jan",1),("fev",2),("mar",3),("avr",4),("mai",5),("juin",6),("juil",7),("aou",8),("sep",9),("oct",10),("nov",11),("dec",12)]
        mo = next((n for k, n in month_map if k in ml), None)
        yy = re.search(r"(\d{2})", s)
        if mo and yy:
            return f"20{yy.group(1)}-{mo:02d}-01"
    return None

def months_between(a: str, b: str) -> int:
    ya, ma, _ = map(int, a.split("-"))
    yb, mb, _ = map(int, b.split("-"))
    return (yb - ya) * 12 + (mb - ma)

def split_addr(addr):
    addr = " ".join(str(addr).split())
    m = re.search(r"^(.*?)\s+(\d{5})$", addr)
    if m:
        cp = m.group(2)
        ville = "Aubagne" if cp == "13400" else ("Marseille" if cp.startswith("13") else "")
        return m.group(1).strip(), cp, ville
    return addr, "", "Marseille"

def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Vision patrimoine"]
    entities_def = [
        {"slug":"beneduc","name":"SCI IR BENEDUC","shortName":"BENEDUC","type":"IR","creation":"","valeurEstimee":380000,"color":"#60a5fa","gradient":"from-blue-500/20 to-transparent","shareholders":[("Johann Faraut",50),("Alexandre Niel",50)],"rows":[8,9,10,11,12,13]},
        {"slug":"troika","name":"SCI IR TROIKA","shortName":"TROIKA","type":"IR","creation":"","valeurEstimee":1265000,"color":"#a78bfa","gradient":"from-violet-500/20 to-transparent","shareholders":[("Johann Faraut",33),("Alexandre Niel",67)],"rows":[22,23,24,25]},
        {"slug":"lavista","name":"SCI IS LA VISTA","shortName":"LA VISTA","type":"IS","creation":"","valeurEstimee":390000,"color":"#22d3ee","gradient":"from-cyan-500/20 to-transparent","shareholders":[("Johann Faraut",33),("Alexandre Niel",67)],"rows":[34,35]},
        {"slug":"rp","name":"Résidence Principale","shortName":"RP","type":"RP","creation":"","valeurEstimee":630000,"color":"#34d399","gradient":"from-emerald-500/20 to-transparent","shareholders":[("Johann Faraut",66),("Alexandre Niel",34)],"rows":[44]},
    ]
    out = {"entities": [], "properties": []}
    for e in entities_def:
        out["entities"].append({k: e[k] for k in ["slug","name","shortName","type","creation","valeurEstimee","color","gradient","shareholders"]})
        debut = AMORT_START[e["slug"]]
        for r in e["rows"]:
            b, c, d = ws[f"B{r}"].value, ws[f"C{r}"].value, ws[f"D{r}"].value
            Ev, Fv, G, H = ws[f"E{r}"].value, ws[f"F{r}"].value, ws[f"G{r}"].value, ws[f"H{r}"].value
            K, L, M = ws[f"K{r}"].value, ws[f"L{r}"].value, ws[f"M{r}"].value
            if not b or str(b).upper() == "TOTAL":
                continue
            address, cp, ville = split_addr(b)
            # Colonne E = CRD à la 1ʳᵉ échéance du tableau d’amort (pas le montant d’achat)
            crd_ref = float(Ev) if isinstance(Ev, (int, float)) else (float(Fv) if isinstance(Fv, (int, float)) else 0)
            crd_proj = float(Fv) if isinstance(Fv, (int, float)) else crd_ref
            prop = {
                "entitySlug": e["slug"], "address": address, "ville": ville or "Marseille", "cp": cp or "13000",
                "type": str(c).strip() if c else "Bien", "lots": int(d) if isinstance(d, (int, float)) else 1,
                "loyer": round(float(G) if isinstance(G, (int, float)) else 0, 2),
                "taxeFonciere": round(float(K) if isinstance(K, (int, float)) else 0, 2),
                "valeurActuelle": round(float(M) if isinstance(M, (int, float)) else 0, 2),
                "prixAchat": 0, "travaux": 0, "fraisNotaire": 0, "assurance": 0,
                "surface": 140 if "140" in str(c) else 0,
            }
            mens = float(H) if isinstance(H, (int, float)) else 0
            if mens > 0 or crd_ref > 0:
                fin = parse_fin(L)
                duree = months_between(debut, fin) if fin else 0
                # capital_restant au seed = vérif Excel à la date de projection par défaut ;
                # l’app recalcule dynamiquement via projectFlatCrd
                prop["credit"] = {
                    "banque": "À préciser",
                    "montantInitial": round(crd_ref, 2),
                    "taux": 0,
                    "duree": max(0, duree),
                    "debut": debut,
                    "mensualite": round(mens, 2),
                    "capitalRestant": round(crd_proj, 2),
                    "finCredit": fin,
                    "amortizationModel": "EXCEL_FLAT",
                }
            out["properties"].append(prop)

    OUT_JSON.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")

    lines = [
        "-- Vision Patrimoine — seed depuis VISION PATRIMOINE 030425.xlsx",
        "-- Régénéré par scripts/excel_to_seed.py",
        "-- Modèle EXCEL_FLAT : CRD = CRD_réf − mensualité × mois (comme Amort Beneduc/Troika)",
        "truncate table amortization_entries, loans, tenants, properties, shareholders, bank_dossiers, alerts, legal_entities cascade;",
        "",
    ]
    for e in out["entities"]:
        eid = uid(f"entity:{e['slug']}")
        lines.append(
            f"insert into legal_entities (id, slug, name, short_name, type, creation, valeur_estimee, color, gradient) values "
            f"('{eid}', '{e['slug']}', $${e['name']}$$, $${e['shortName']}$$, '{e['type']}', null, {e['valeurEstimee']}, '{e['color']}', '{e['gradient']}');"
        )
        for name, parts in e["shareholders"]:
            lines.append(
                f"insert into shareholders (id, entity_id, name, parts) values ('{uid(f'sh:{e['slug']}:{name}')}', '{eid}', $${name}$$, {parts});"
            )
        lines.append("")
    for p in out["properties"]:
        eid = uid(f"entity:{p['entitySlug']}")
        pid = uid(f"prop:{p['entitySlug']}:{p['address']}")
        lines.append(
            f"insert into properties (id, entity_id, address, ville, cp, type, surface, lots, prix_achat, travaux, frais_notaire, valeur_actuelle, loyer, taxe_fonciere, assurance) values "
            f"('{pid}', '{eid}', $${p['address']}$$, $${p['ville']}$$, '{p['cp']}', $${p['type']}$$, {p['surface']}, {p['lots']}, "
            f"{p['prixAchat']}, {p['travaux']}, {p['fraisNotaire']}, {p['valeurActuelle']}, {p['loyer']}, {p['taxeFonciere']}, {p['assurance']});"
        )
        c = p.get("credit")
        if c:
            fin = f"'{c['finCredit']}'::date" if c.get("finCredit") else "null"
            debut = f"'{c['debut']}'::date" if c.get("debut") else "null"
            model = c.get("amortizationModel", "EXCEL_FLAT")
            lines.append(
                f"insert into loans (id, property_id, banque, montant_initial, taux_annuel, duree_mois, date_debut, assurance_mensuelle, mensualite, capital_restant, fin_credit, amortization_model) values "
                f"('{uid(f'loan:{p['entitySlug']}:{p['address']}')}', '{pid}', $${c['banque']}$$, {c['montantInitial']}, {c['taux']}, {c['duree']}, {debut}, 0, {c['mensualite']}, {c['capitalRestant']}, {fin}, '{model}');"
            )
        lines.append("")
    lines.append(
        f"insert into alerts (id, type, title, detail, severity) values ('{uid('alert:1')}', 'info', 'Données importées depuis Excel', 'Patrimoine Vision — amortissement flat aligné Excel (CRD − mensualité).', 'low');"
    )
    OUT_SQL.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"OK → {OUT_JSON.name}, {OUT_SQL.name} ({len(out['properties'])} biens)")

if __name__ == "__main__":
    main()
